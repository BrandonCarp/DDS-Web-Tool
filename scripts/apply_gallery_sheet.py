#!/usr/bin/env python3
"""
Load the Gallery GD1SP/GD1LP pricing sheet into the app.

    python scripts/apply_gallery_sheet.py "PRICING 10-8.xlsx" [--margin 43]

The sheet (one tab, "GD1LPGD1SP") carries DDS COST, so:

  * the base door grid is written into src/lib/pricing/data/special-doors.ts
    as SELL — cost / (1 - margin), the Gallery door margin — because that grid
    holds sell figures for every model (see the header of that file);
  * the glass tables and Ultra-Grain adders are written into
    src/lib/pricing/data/gallery-glass.ts as COST, because the special order
    quote lifts adders by the margin itself at quote time, the way the 4050's
    glass tables work.

Sheet layout it reads:
  rows 4-5      ULTRA GRAIN: "SINGLE $x  DOUBLE $y  UP TO 8FT HIGH" and
                "... 8FT AND UP" (8'2" and up — Brandon, 8/10/2026)
  row 8         height headings: 7FT | 8FT | 9FT AND 10FT, in D, F, H
  rows 9-28     width band in A, solid door cost under each heading
  row 31 on     glass blocks: a type name in C (and its /INSERTS twin in E),
                then a GD1SP half and a GD1LP half, each "band | windows |
                glass | with inserts" per row. ACRYLIC has no inserts column.

Checks before writing: every window block must be its per-window figure
times the window count, and the inserts column the glass column plus the
same flat amount per window — the shape the sheet had on 8/10/2026. A cell
that breaks the pattern is reported and the script stops, so a typo never
reaches the app silently (E192 was 103.11 for 824.90 on the first copy).
"""
import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SPECIAL_DOORS = ROOT / "src/lib/pricing/data/special-doors.ts"
GALLERY_GLASS = ROOT / "src/lib/pricing/data/gallery-glass.ts"
GROUP = "GD1LP/GD1SP"

# Every width the app grids, 6'0" to 18'0" by 2": the band a width falls in
# gives its price.
WIDTH_KEYS = [f"{ft}" if inch == 0 else f"{ft}.{inch}"
              for ft in range(6, 19) for inch in (0, 2, 4, 6, 8, 10) if not (ft == 18 and inch > 0)]


def inches(label: str) -> int:
    m = re.match(r"\s*(\d+)'\s*(\d+)?", label)
    return int(m.group(1)) * 12 + int(m.group(2) or 0)


def band_range(label: str) -> tuple[int, int]:
    """'8\'2" to 8\'10"' -> (98, 106); '8\'0"' -> (96, 96); 'UP TO 7\'6"' -> (0, 90)."""
    s = label.strip().upper().replace("TO", " TO ")
    if s.startswith("UP"):
        return 0, inches(s.split("TO", 1)[1])
    parts = [p for p in s.split(" TO ") if p.strip()]
    lo = inches(parts[0])
    hi = inches(parts[1]) if len(parts) > 1 else lo
    return lo, hi


def width_inches(key: str) -> int:
    ft, _, inch = key.partition(".")
    return int(ft) * 12 + int(inch or 0)


def money(x: float) -> float:
    return round(x + 1e-9, 2)


def read_sheet(path: Path):
    ws = load_workbook(path, data_only=True).active
    # Ultra-Grain
    ug = {}
    for r in (4, 5):
        text = str(ws[f"B{r}"].value)
        nums = [float(x.replace(",", "")) for x in re.findall(r"\$\s*([\d,]+\.\d+)", text)]
        if len(nums) != 2:
            sys.exit(f"could not read the Ultra-Grain line in B{r}: {text!r}")
        key = "low" if "UP TO" in text.upper() else "high"
        ug[key] = {"single": nums[0], "double": nums[1]}
    if set(ug) != {"low", "high"}:
        sys.exit("expected one 'UP TO 8FT HIGH' and one '8FT AND UP' Ultra-Grain line")

    # Base grid: A = band, D/F/H = 7 / 8 / 9-and-10 ft
    cols = {"7": "D", "8": "F", "9": "H", "10": "H"}
    heads = {ws["D8"].value, ws["F8"].value, ws["H8"].value}
    if heads != {"7FT", "8FT", "9FT AND 10FT"}:
        sys.exit(f"row 8 headings changed: {heads}")
    bands = []
    for r in range(9, 29):
        label = ws[f"A{r}"].value
        if not label:
            break
        lo, hi = band_range(label)
        bands.append((lo, hi, {t: float(ws[f"{c}{r}"].value) for t, c in cols.items()}))

    def base_cost(width_key: str, tier: str) -> float | None:
        w = width_inches(width_key)
        for lo, hi, prices in bands:
            if lo <= w <= hi:
                return prices[tier]
        return None

    # Glass blocks
    blocks = []
    r = 30
    while r <= ws.max_row:
        name = ws[f"C{r}"].value
        if isinstance(name, str) and ws[f"B{r}"].value is None and ws[f"A{r}"].value is None:
            blocks.append((r, name.strip()))
        r += 1
    glass = {}
    problems = []
    for i, (r0, name) in enumerate(blocks):
        end = blocks[i + 1][0] if i + 1 < len(blocks) else ws.max_row + 1
        model = None
        per_window = None
        per_insert = None
        table = {"GD1SP": [], "GD1LP": []}
        for rr in range(r0 + 1, end):
            b = ws[f"B{rr}"].value
            if b in ("GD1SP", "GD1LP"):
                model = b
                continue
            if model and isinstance(b, (int, float)) and ws[f"A{rr}"].value:
                n = int(b)
                g = float(ws[f"C{rr}"].value)
                gi = ws[f"E{rr}"].value
                gi = float(gi) if gi is not None else None
                units = n if model == "GD1SP" else n * 2   # a long window is two short ones
                pw = g / units
                if per_window is None:
                    per_window = pw
                elif abs(pw - per_window) > 0.02:
                    problems.append(f"{name}: {model} {ws[f'A{rr}'].value} glass {g} is {pw:.2f}/window, block runs {per_window:.2f}")
                if gi is not None:
                    pi = (gi - g) / units
                    if per_insert is None:
                        per_insert = pi
                    elif abs(pi - per_insert) > 0.02:
                        problems.append(f"{name}: {model} {ws[f'A{rr}'].value} inserts {gi} adds {pi:.2f}/window, block runs {per_insert:.2f} (cell E{rr})")
                lo, hi = band_range(ws[f"A{rr}"].value)
                table[model].append({"lo": lo, "hi": hi, "windows": n, "glass": g, "inserts": gi})
        glass[name] = table
    if problems:
        print("The sheet disagrees with itself — fix these before loading:", file=sys.stderr)
        for p in problems:
            print("  -", p, file=sys.stderr)
        sys.exit(1)
    return ug, base_cost, glass


# Sheet names -> app ids and labels, in the order the counter reads them.
GLASS_IDS = [
    ("DSB", "dsb", "Double strength"),
    ("ACRYLIC", "acrylic", "Acrylic"),
    ("OBSCURE", "obscure", "Obscure"),
    ("INSULATED", "insulated", "Insulated"),
    ("INS OBSCURE GLASS", "insulated_obscure", "Insulated obscure"),
    ('1/8" FROSTED DSB', "frosted", "Frosted DSB"),
    ("INS FROSTED GLASS", "insulated_frosted", "Insulated frosted"),
    ('1/8" RAIN DSB', "rain", "Rain DSB"),
    ("INS RAIN GLASS", "insulated_rain", "Insulated rain"),
    ('1/8" SEEDED DSB', "seeded", "Seeded DSB"),
    ("INS SEEDED GLASS", "insulated_seeded", "Insulated seeded"),
    ('1/8" MIDNIGHT GRAY DSB', "midnight_gray", "Midnight gray DSB"),
    ("INS MIDNIGHT GRAY GLASS", "insulated_midnight_gray", "Insulated midnight gray"),
    ('1/8" TEMP MIDNIGHT GRAY', "tempered_midnight_gray", "Tempered midnight gray"),
    ("INS MIDNIGHT GRAY TEMPERED", "insulated_tempered_midnight_gray", "Insulated tempered midnight gray"),
]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--margin", type=float, default=43.0, help="Gallery door margin, percent (default 43)")
    args = ap.parse_args()
    ug, base_cost, glass = read_sheet(Path(args.workbook))
    missing = [s for s, _, _ in GLASS_IDS if s not in glass]
    extra = [s for s in glass if s not in {g[0] for g in GLASS_IDS}]
    if missing or extra:
        sys.exit(f"glass blocks changed — missing {missing}, unexpected {extra}")
    sell = lambda cost: money(cost / (1 - args.margin / 100))

    # ---- the base grid, as SELL. DSB is the Gallery's base glass: GLASS and
    # INSERTS are the solid door plus the DSB adder for that width and model.
    # GD1SP and GD1LP share one door price; their glass differs, so the grid
    # carries the GD1SP (short) figure and the quote re-prices glass by model.
    dsb = glass["DSB"]["GD1SP"]

    def dsb_add(width_key: str, inserts: bool) -> float | None:
        w = width_inches(width_key)
        for row in dsb:
            if row["lo"] <= w <= row["hi"]:
                return row["inserts"] if inserts else row["glass"]
        return None

    grid = {}
    for tier in ("7", "8", "9", "10"):
        grid[tier] = {}
        for wk in WIDTH_KEYS:
            c = base_cost(wk, tier)
            g = dsb_add(wk, False)
            gi = dsb_add(wk, True)
            if c is None or g is None or gi is None:
                continue
            grid[tier][wk] = {"solid": sell(c), "glass": sell(c + g), "inserts": sell(c + gi)}

    src = SPECIAL_DOORS.read_text()
    # Replace an existing Gallery block, or add one before the closing of SPECIAL_DOORS.
    block = f'  "{GROUP}": ' + json.dumps(grid, indent=2).replace("\n", "\n  ")
    pat = re.compile(r'  "' + re.escape(GROUP) + r'": \{.*?\n  \}(,?)\n', re.S)
    if pat.search(src):
        src = pat.sub(lambda m: block + m.group(1) + "\n", src, count=1)
    else:
        i = src.rindex("\n};")
        src = src[:i].rstrip() + ",\n" + block + "\n};" + src[i + 3:]
    SPECIAL_DOORS.write_text(src)

    # ---- the glass tables and Ultra-Grain, as COST
    out = ["// AUTO-GENERATED by scripts/apply_gallery_sheet.py from the Gallery",
           "// GD1SP/GD1LP pricing sheet (PRICING 10-8.xlsx, 8/10/2026). Do not edit by",
           "// hand — re-run the script against a newer sheet.",
           "//",
           "// Everything here is DDS COST, before margin. The special order quote lifts",
           "// it by the Gallery door margin at quote time, the way the 4050's glass",
           "// tables in so-glass.ts are handled. Width bands are in inches.",
           "",
           "export interface GalleryBand {",
           "  /** Door width this band covers, in inches, inclusive. */",
           "  lo: number;",
           "  hi: number;",
           "  /** Windows across the door in this band. */",
           "  windows: number;",
           "  /** Glass, and glass with decorative inserts; null where inserts are not offered. */",
           "  glass: number;",
           "  inserts: number | null;",
           "}",
           "",
           "export const GALLERY_GLASS_TYPES: { id: string; label: string }[] = [",
           *[f'  {{ id: "{gid}", label: "{label}" }},' for _, gid, label in GLASS_IDS],
           "];",
           "",
           "/** glass id -> model -> bands, narrowest first. */",
           "export const GALLERY_GLASS: Record<string, Record<string, GalleryBand[]>> = {"]
    for sheet_name, gid, _ in GLASS_IDS:
        out.append(f'  "{gid}": {{')
        for model in ("GD1SP", "GD1LP"):
            rows = glass[sheet_name][model]
            out.append(f'    "{model}": [')
            for row in rows:
                ins = "null" if row["inserts"] is None else f"{row['inserts']:.2f}"
                out.append(f'      {{ lo: {row["lo"]}, hi: {row["hi"]}, windows: {row["windows"]}, glass: {row["glass"]:.2f}, inserts: {ins} }},')
            out.append("    ],")
        out.append("  },")
    out += ["};",
            "",
            "/**",
            " * Ultra-Grain (and the other premium finishes) on a Gallery door: a flat",
            " * adder by door, single or double, in two height bands — up to 8'0\", and",
            " * 8'2\" and up (Brandon, 8/10/2026). Cost, before margin.",
            " */",
            "export const GALLERY_ULTRA_GRAIN = {",
            f'  upTo8ft: {{ single: {ug["low"]["single"]:.2f}, double: {ug["low"]["double"]:.2f} }},',
            f'  over8ft: {{ single: {ug["high"]["single"]:.2f}, double: {ug["high"]["double"]:.2f} }},',
            "} as const;",
            ""]
    GALLERY_GLASS.write_text("\n".join(out))
    n = sum(len(t) for t in grid.values())
    print(f"wrote {n} grid cells for {GROUP} at {args.margin:g}% into {SPECIAL_DOORS.name}")
    print(f"wrote {len(GLASS_IDS)} glass types and Ultra-Grain into {GALLERY_GLASS.name}")


if __name__ == "__main__":
    main()
