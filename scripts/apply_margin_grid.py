"""
Special order grids for models priced at a margin — Brandon, 2/10/2026, for
the 4308 and 4138.

    python scripts/apply_margin_grid.py WORKBOOK.xlsx --margin 45 4308 4138
    python scripts/apply_margin_grid.py WORKBOOK.xlsx --margin 45 --hold 4308:7:8.6:glass 4308 4138

Each model's tabs are found by name ("4308 7ft", "4308 8ft"...), and the
height comes from the tab's name, so a typo in a row's height cannot drop a
price (such rows are reported). The price used is the tab's SELL, rounded to
the cent like Excel's ROUND; where SELL is blank it is TOTAL / (1 - margin),
the same thing. Any row where SELL and TOTAL disagree at that margin is
reported. A size marked N/A is left off, so it asks for the typed total.

--hold MODEL:HEIGHT:WIDTH:STYLE keeps one cell off the grid on purpose — a
price that looks wrong on the sheet asks for the typed total until the sheet
is put right and this is re-run without the hold.

Re-running replaces those models' grids entirely, so sizes added to the sheet
later simply appear.
"""
import argparse
import json
import re
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
TS = ROOT / "src/lib/pricing/data/special-doors.ts"
WIDTH = re.compile(r"^(\d+)'(\d+)\"?\s*X\s*(.*)$")
HEIGHT = re.compile(r"^(\d+)'(\d+)\"?$")
STYLES = ("solid", "glass", "inserts")


def width_key(feet: int, inches: int) -> str:
    return str(feet) if inches == 0 else f"{feet}.{inches}"


def cents(x) -> float:
    return float(Decimal(str(x)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def read_tab(ws, height: str, margin: int):
    """{width: {style: sell}} for the tab's one height."""
    grid = {}
    for n, row in enumerate(ws.iter_rows(values_only=True), start=1):
        size = str(row[0] or "").strip().upper().replace("  ", " ")
        m = WIDTH.match(size)
        style = str(row[1] or "").strip().lower()
        if not m or style not in STYLES or not isinstance(row[4], (int, float)):
            continue
        wf, wi, written = int(m.group(1)), int(m.group(2)), m.group(3).strip()
        h = HEIGHT.match(written)
        if not h or (int(h.group(1)), int(h.group(2))) != (int(height), 0):
            print(f"    row {n}: height written as {written!r}, read as the tab's {height}'0\"")
        from_total = cents(Decimal(str(row[4])) / (Decimal(1) - Decimal(margin) / 100))
        sell = cents(row[6]) if isinstance(row[6], (int, float)) else from_total
        if sell != from_total:
            print(f"    row {n}: {size} {style} SELL {sell} is not TOTAL at {margin}M ({from_total}) — SELL used")
        grid.setdefault(width_key(wf, wi), {})[style] = sell
    return grid


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook")
    ap.add_argument("--margin", type=int, required=True)
    ap.add_argument("--hold", action="append", default=[], metavar="MODEL:HEIGHT:WIDTH:STYLE")
    ap.add_argument("models", nargs="+")
    args = ap.parse_args()

    wb = load_workbook(args.workbook, data_only=True)
    src = TS.read_text()
    start = src.index("{", src.index("export const SPECIAL_DOORS"))
    payload = json.loads(src[start:].rsplit(";", 1)[0])
    for model in args.models:
        tabs = [n for n in wb.sheetnames if re.fullmatch(rf"{re.escape(model)}\s+\d+ft\s*", n)]
        if not tabs:
            raise SystemExit(f"ABORT — no '{model} <height>ft' tabs in the workbook")
        payload[model] = {}
        for name in tabs:
            height = re.search(r"(\d+)ft", name).group(1)
            payload[model][height] = read_tab(wb[name], height, args.margin)
            print(f"  {name.strip():10s} {len(payload[model][height])} widths priced at {args.margin}M")
    for hold in args.hold:
        model, height, width, style = hold.split(":")
        cell = payload.get(model, {}).get(height, {}).get(width, {})
        if style not in cell:
            raise SystemExit(f"ABORT — nothing to hold at {hold}")
        print(f"  held off the grid: {model} {width} x {height}ft {style} (was {cell.pop(style)}) — asks for the typed total")
    TS.write_text(src[:start] + json.dumps(payload, indent=2) + ";\n")
    print(f"wrote {TS.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
