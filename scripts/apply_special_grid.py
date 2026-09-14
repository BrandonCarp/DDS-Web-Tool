#!/usr/bin/env python3
"""
Add or refresh a special-order grid in special-doors.ts from workbook tabs.

Different from apply_pricing_workbook.py, which rewrites the residential and
stock grids. This one touches only SPECIAL_DOORS — the 73-or-so width tables the
special order configurator reads.

Two things it normalises, both of which would break lookups silently:

  * 11-foot widths are written 11.00, 11.02 ... 11.10 on some sheets and
    11, 11.2 ... 11.10 on others. They key as the second form everywhere.
  * a tab name may order the models differently from the margin group key
    ("4300-4310-4301" vs "4300/4301/4310"), so the key is given explicitly.

Margins are verified, never corrected: SELL is the column DDS quotes, so a row
that does not reproduce from TOTAL is reported and taken as written.

Usage:
  python3 scripts/apply_special_grid.py <workbook.xlsx> "<group key>" <tier>=<tab> ...

Example:
  python3 scripts/apply_special_grid.py book.xlsx "4300/4301/4310" \
      7="4300-4310-4301 7ft" 8="4300-4310-4301 8ft"
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
TARGET = ROOT / "src/lib/pricing/data/special-doors.ts"
SIZE = re.compile(r"^(\d+)'(\d+)\"?\s*X\s*(\d+)'(\d+)\"?$", re.I)
STYLES = ("solid", "glass", "inserts")


def width_key(ft: str, inch: str) -> str:
    """6'2" -> "6.2", 11'0" -> "11". Strips the leading zero some sheets write."""
    n = int(inch)
    return ft if n == 0 else f"{ft}.{n}"


def read_tab(ws, margin: float, label: str):
    grid: dict[str, dict[str, float]] = {}
    off: list[str] = []
    for row in ws.iter_rows(values_only=True):
        head = str(row[0]).strip().upper().replace("  ", " ") if row[0] else ""
        m = SIZE.match(head)
        if not m or len(row) < 7:
            continue
        total, sell = row[4], row[6]
        if not isinstance(sell, (int, float)) or not isinstance(total, (int, float)):
            continue
        style = str(row[1]).strip().lower()
        if style not in STYLES:
            continue
        expect = total / (1 - margin / 100)
        if abs(expect - sell) > 1.00:
            off.append(f"{head} {style}: {sell:.2f} vs {expect:.2f} ({sell - expect:+.2f})")
        grid.setdefault(width_key(m.group(1), m.group(2)), {})[style] = round(float(sell), 2)

    missing = {w: [s for s in STYLES if s not in d] for w, d in grid.items()}
    missing = {w: v for w, v in missing.items() if v}
    if missing:
        raise SystemExit(f"  {label}: {len(missing)} width(s) missing a style — {list(missing)[:4]}")
    return grid, off


def marker_of(ws) -> int | None:
    for row in ws.iter_rows(values_only=True):
        for c in row:
            if isinstance(c, str) and re.fullmatch(r"\d+M", c.strip()):
                return int(c.strip()[:-1])
    return None


def main():
    if len(sys.argv) < 4:
        sys.exit(__doc__)
    book, key, pairs = sys.argv[1], sys.argv[2], sys.argv[3:]
    wb = load_workbook(book, data_only=True)

    tiers: dict[str, dict] = {}
    total_off = 0
    for pair in pairs:
        tier, tab = pair.split("=", 1)
        match = next((s for s in wb.sheetnames if s.strip() == tab.strip()), None)
        if not match:
            sys.exit(f"  no tab named {tab!r}; have {wb.sheetnames}")
        ws = wb[match]
        margin = marker_of(ws)
        if margin is None:
            sys.exit(f"  {tab}: no NNM margin marker found")
        grid, off = read_tab(ws, margin, tab)
        total_off += len(off)
        print(f"  {tab.strip():26s} {margin}M  {len(grid)} widths  {sum(len(d) for d in grid.values())} prices"
              f"  off-margin {len(off)}")
        for o in off[:6]:
            print(f"      {o}")
        tiers[tier] = grid

    src = TARGET.read_text(encoding="utf-8")
    body = src[src.index("{", src.index("SPECIAL_DOORS")):src.rindex(";")]
    data = json.loads(body)
    existing = key in data
    data[key] = {t: tiers[t] for t in sorted(tiers, key=int)}

    out = json.dumps(data, indent=2)
    src = src[:src.index("{", src.index("SPECIAL_DOORS"))] + out + src[src.rindex(";"):]
    TARGET.write_text(src, encoding="utf-8")
    print(f"\n  {'refreshed' if existing else 'added'} {key}: tiers {sorted(tiers, key=int)}")
    print(f"  rows off their stated margin (taken as written): {total_off}")


if __name__ == "__main__":
    main()
